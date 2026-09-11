#!/usr/bin/env python3
"""Informes DOCX/PDF del curso, exclusivamente desde contenido y evidencia aportados.

Uso: python scripts/reportes_curso.py datos.json --salida informes/version-nueva
Consultar docs/plantilla/CONTRATO.md. Nunca sobrescribe archivos existentes.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from datetime import datetime, date
from pathlib import Path
from urllib.parse import urlparse
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from PIL import Image as PILImage
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Image, KeepTogether,
    Table, TableStyle, HRFlowable,
)

AZUL = "12235C"
GRIS = "666666"
ANCHO = 612 - 2 * 52.8
ALTO_IMAGEN = 350
CURSO = "Programación de Backend y MCP en Python para IA Generativa"
ESTADOS = {"final", "borrador", "prueba", "entrega_con_observaciones"}
ESTADOS_REVISADOS = {"final", "entrega_con_observaciones"}
AVISOS_ESTADO = {
    "borrador": "BORRADOR PARA REVISIÓN",
    "prueba": "PRUEBA TÉCNICA DEL GENERADOR - NO ENTREGAR",
    "entrega_con_observaciones": "INFORME DE ENTREGA CON OBSERVACIONES",
}
TIPOS = {"captura", "registro_renderizado", "simulacion"}
PRACTICAS_CON_REPOSITORIO = {1, 3, 4}
EVIDENCIAS_OBLIGATORIAS = {
    1: {"conversacion_8_turnos", "manejo_429", "repositorio"},
    2: {"cierre"},
    3: {"spec_manual", "flujo_speckit"},
    4: {"hook_bloqueando", "reporte_qa"},
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def texto(value: object, campo: str, *, opcional: bool = False) -> str:
    if not isinstance(value, str) or (not value.strip() and not opcional):
        raise ValueError(f"{campo}: se necesita texto no vacío")
    return value.strip()


def parrafos(value: object, campo: str) -> list[str]:
    if isinstance(value, str):
        return [p.strip() for p in value.split("\n\n") if p.strip()]
    if isinstance(value, list):
        return [texto(p, campo) for p in value]
    raise ValueError(f"{campo}: debe ser texto o lista de párrafos")


def archivo_relativo(base: Path, value: object, campo: str) -> Path:
    p = Path(texto(value, campo)).expanduser()
    p = (base / p).resolve() if not p.is_absolute() else p.resolve()
    if not p.is_file():
        raise ValueError(f"{campo}: no existe un archivo legible: {p.name}")
    return p


def validar_observaciones(datos: dict, ids: set[str]) -> set[str]:
    """Admite únicamente el 429 no observado de P1, con respaldo identificable."""
    observaciones = datos.setdefault("observaciones", [])
    if datos["estado"] != "entrega_con_observaciones":
        if observaciones:
            raise ValueError("observaciones solo se admite en estado entrega_con_observaciones")
        return set()
    if datos["practica"]["numero"] != 1:
        raise ValueError("entrega_con_observaciones solo está habilitado para práctica 1")
    if not isinstance(observaciones, list) or len(observaciones) != 1 or not isinstance(observaciones[0], dict):
        raise ValueError("La entrega necesita una observación explícita de manejo_429")
    obs = observaciones[0]
    if obs.get("requisito_id") != "manejo_429" or obs.get("tipo") != "no_observado":
        raise ValueError("Solo se admite la observación manejo_429 de tipo no_observado")
    if "manejo_429" in ids:
        raise ValueError("La evidencia manejo_429 contradice la observación no_observado")
    referencias = obs.get("evidencia_ids")
    if (not isinstance(referencias, list) or not referencias
            or any(not isinstance(eid, str) for eid in referencias)
            or len(set(referencias)) != len(referencias)
            or "ensayo_sin_429" not in referencias or set(referencias) - ids):
        raise ValueError("La observación necesita ensayo_sin_429 y referencias a evidencias existentes")
    for campo in ("protocolo", "resultado", "limite"):
        obs[campo] = texto(obs.get(campo), f"observaciones.{campo}")
    return {"manejo_429"}


def cargar(path: Path) -> dict:
    datos = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(datos, dict) or datos.get("version") != 1:
        raise ValueError("version debe ser 1 y la raíz debe ser un objeto JSON")
    estado = datos.get("estado", "borrador")
    if estado not in ESTADOS:
        raise ValueError("estado debe ser final, borrador, prueba o entrega_con_observaciones")
    datos["estado"] = estado
    practica = datos.get("practica")
    if not isinstance(practica, dict) or practica.get("numero") not in (1, 2, 3, 4):
        raise ValueError("practica.numero debe estar entre 1 y 4")
    practica["titulo"] = texto(practica.get("titulo"), "practica.titulo")
    practica["clase"] = texto(practica.get("clase", ""), "practica.clase", opcional=True)
    estudiante = datos.get("estudiante")
    if not isinstance(estudiante, dict):
        raise ValueError("estudiante debe incluir nombre")
    estudiante["nombre"] = texto(estudiante.get("nombre"), "estudiante.nombre")
    hoy = datetime.now(ZoneInfo("America/Guayaquil")).date()
    fecha = datos.get("fecha", hoy.isoformat())
    if date.fromisoformat(fecha) > hoy:
        raise ValueError("fecha no puede estar en el futuro")
    datos["fecha"] = fecha
    datos["explicacion"] = parrafos(datos.get("explicacion", ""), "explicacion")
    if not datos["explicacion"]:
        raise ValueError("explicacion no puede estar vacía")
    datos["reflexion"] = parrafos(datos.get("reflexion", []), "reflexion")
    for bloque in datos.get("secciones", []):
        bloque["titulo"] = texto(bloque.get("titulo"), "secciones.titulo")
        bloque["parrafos"] = parrafos(bloque.get("parrafos", []), "secciones.parrafos")
        if bloque.get("posicion", "despues_evidencia") not in ("antes_evidencia", "despues_evidencia"):
            raise ValueError("secciones.posicion no reconocida")
        if "tabla" in bloque:
            tabla = bloque["tabla"]
            columnas = tabla.get("columnas", [])
            if not 1 <= len(columnas) <= 6:
                raise ValueError("Cada tabla necesita de 1 a 6 columnas")
            tabla["columnas"] = [texto(x, "tabla.columnas") for x in columnas]
            if not tabla.get("filas"):
                raise ValueError("La tabla debe contener filas")
            for fila in tabla["filas"]:
                if len(fila) != len(columnas):
                    raise ValueError("Las filas de tabla deben tener todas las columnas")
                for celda in fila:
                    texto(celda, "tabla.filas", opcional=True)
            anchos = tabla.get("anchos", [1] * len(columnas))
            if len(anchos) != len(columnas) or any(not isinstance(n, (int, float)) or n <= 0 for n in anchos):
                raise ValueError("tabla.anchos necesita un peso positivo por columna")
            tabla["anchos"] = [n / sum(anchos) * ANCHO for n in anchos]
    evidencias = datos.setdefault("evidencias", [])
    ids = set()
    for evidencia in evidencias:
        eid = texto(evidencia.get("id"), "evidencias.id")
        if not re.fullmatch(r"[a-z0-9_]+", eid) or eid in ids:
            raise ValueError("Cada evidencia necesita id único con letras minúsculas, números o _")
        ids.add(eid)
        evidencia["titulo"] = texto(evidencia.get("titulo"), "evidencias.titulo")
        evidencia["descripcion"] = texto(evidencia.get("descripcion"), "evidencias.descripcion")
        tipo = evidencia.get("tipo")
        if tipo not in TIPOS:
            raise ValueError(f"Evidencia {eid}: tipo debe ser captura, registro_renderizado o simulacion")
        imagen = archivo_relativo(path.parent, evidencia.get("imagen"), "evidencias.imagen")
        with PILImage.open(imagen) as img:
            img.verify()
        evidencia["_imagen"] = imagen
        fuente = archivo_relativo(path.parent, evidencia.get("fuente"), "evidencias.fuente")
        evidencia["_fuente"] = fuente
        evidencia["_sha256_imagen"] = sha256(imagen)
        evidencia["_sha256_fuente"] = sha256(fuente)
        for clave, real in (("sha256_imagen", evidencia["_sha256_imagen"]), ("sha256_fuente", evidencia["_sha256_fuente"])):
            if clave in evidencia and evidencia[clave] != real:
                raise ValueError(f"Evidencia {eid}: {clave} no coincide")
        if estado in ESTADOS_REVISADOS and (tipo == "simulacion" or evidencia.get("verificada") is not True):
            raise ValueError(f"Evidencia {eid}: una entrega revisada exige evidencia verificada, sin simulación")
    pendientes_permitidos = validar_observaciones(datos, ids)
    repo = datos.setdefault("repositorio", {})
    url = repo.get("url", "")
    if url:
        parsed = urlparse(url)
        if parsed.scheme != "https" or not parsed.netloc or parsed.username or parsed.password or parsed.query:
            raise ValueError("repositorio.url debe ser HTTPS, sin credenciales ni parámetros")
        if parsed.hostname in ("example.com", "example.org", "example.net") and estado in ESTADOS_REVISADOS:
            raise ValueError("Una entrega revisada no admite un repositorio de ejemplo")
    pendientes = []
    for requisito in datos.get("requisitos", []):
        texto(requisito.get("descripcion"), "requisitos.descripcion")
        if requisito.get("estado") not in ("cumplido", "pendiente", "no_aplica"):
            raise ValueError("requisitos.estado no reconocido")
        if set(requisito.get("evidencia_ids", [])) - ids:
            raise ValueError("Un requisito refiere a evidencias inexistentes")
        if requisito["estado"] == "pendiente":
            pendientes.append(requisito)
            if estado in ESTADOS_REVISADOS and requisito.get("id") not in pendientes_permitidos:
                raise ValueError("Un informe final no admite requisitos pendientes salvo la observación documentada de manejo_429")
    if estado == "entrega_con_observaciones":
        relacionados = [r for r in datos.get("requisitos", []) if r.get("id") == "manejo_429"]
        if (len(pendientes) != 1 or len(relacionados) != 1 or relacionados[0]["estado"] != "pendiente"
                or not set(datos["observaciones"][0]["evidencia_ids"]).issubset(relacionados[0].get("evidencia_ids", []))):
            raise ValueError("Declarar una sola vez el requisito manejo_429 como pendiente y vincular su evidencia de ensayo")
    if estado in ESTADOS_REVISADOS:
        if not evidencias:
            raise ValueError("Un informe final necesita evidencia")
        numero = practica["numero"]
        if (numero in PRACTICAS_CON_REPOSITORIO or url) and (
            not url or repo.get("acceso_docente_verificado") is not True
        ):
            raise ValueError(f"Práctica {numero}: el repositorio incluido o exigido necesita URL y acceso docente verificado")
        faltan = EVIDENCIAS_OBLIGATORIAS[numero] - ids - pendientes_permitidos
        if faltan:
            raise ValueError(f"Práctica {numero}: faltan las evidencias exigidas: {', '.join(sorted(faltan))}")
        if numero == 2:
            if len(ids & {"ronda_1", "ronda_2", "ronda_3"}) < 2:
                raise ValueError("Práctica 2: se necesita evidencia de al menos dos rondas distintas")
            if len(datos["reflexion"]) != 3:
                raise ValueError("Práctica 2: se necesitan las tres respuestas de reflexión, una por párrafo")
        if numero == 3 and not any("tabla" in bloque for bloque in datos.get("secciones", [])):
            raise ValueError("Práctica 3: se necesita la tabla comparativa en secciones")
        if datos.get("contenido_revisado") is not True:
            raise ValueError("Marcar contenido_revisado=true después de verificar el contenido y su autoría")
    datos["_entrada"] = path
    return datos


def subtitulo(d: dict) -> str:
    p = d["practica"]
    clase = f" - Clase {p['clase']}" if p["clase"] else ""
    return f"Práctica {p['numero']}{clase}: {p['titulo']} | {CURSO}"


def mostrar_repositorio(d: dict) -> bool:
    """La guía de práctica 2 excluye expresamente la entrega de proyecto."""
    return d["practica"]["numero"] in PRACTICAS_CON_REPOSITORIO or bool(d["repositorio"].get("url"))


def etiqueta_evidencia(e: dict, numero: int) -> str:
    tipo = {"captura": "Captura", "registro_renderizado": "Registro de ejecución", "simulacion": "Simulación de prueba"}[e["tipo"]]
    return f"{tipo} {numero} - {e['titulo']}"


def parrafos_observaciones(d: dict) -> list[str]:
    """Texto común para que PDF y DOCX conserven el mismo límite de evidencia."""
    if d["estado"] != "entrega_con_observaciones":
        return []
    obs = d["observaciones"][0]
    return [
        "Este informe documenta el experimento realizado y un requisito pendiente. No declara cumplimiento total de los requisitos ni acredita la observación de un error 429 real.",
        "Requisito manejo_429: pendiente. Tipo de observación: no_observado.",
        f"Protocolo ejecutado: {obs['protocolo']}",
        f"Resultado observado: {obs['resultado']}",
        f"Límite de la conclusión: {obs['limite']}",
        f"Evidencias de respaldo: {', '.join(obs['evidencia_ids'])}.",
    ]


def dimensiones(path: Path) -> tuple[float, float]:
    with PILImage.open(path) as im:
        w, h = im.size
    escala = min(ANCHO / w, ALTO_IMAGEN / h)
    return w * escala, h * escala


def pdf_bytes(d: dict) -> bytes:
    salida = io.BytesIO()
    styles = {
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=7),
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#" + AZUL), spaceAfter=2),
        "sub": ParagraphStyle("sub", fontName="Helvetica", fontSize=11, leading=13, textColor=colors.HexColor("#" + GRIS), spaceAfter=12),
        "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor("#" + AZUL), spaceBefore=12, spaceAfter=8, keepWithNext=True),
        "caption": ParagraphStyle("caption", fontName="Helvetica-Oblique", fontSize=9, leading=12, textColor=colors.HexColor("#" + GRIS), spaceBefore=7, spaceAfter=5),
        "small": ParagraphStyle("small", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#" + GRIS), spaceAfter=7),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=9, leading=12),
        "cell_h": ParagraphStyle("cell_h", fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=colors.white),
    }

    def p(s, style="body"):
        return Paragraph(escape(str(s)).replace("\n", "<br/>"), styles[style])

    story = [p("Informe de Entrega", "title"), p(subtitulo(d), "sub"),
             HRFlowable(width="100%", thickness=1, color=colors.HexColor("#" + AZUL)), Spacer(1, 5)]
    if d["estado"] in AVISOS_ESTADO:
        story.append(p(AVISOS_ESTADO[d["estado"]], "h"))
    story += [p("Datos del estudiante", "h"), p(f"Nombre completo: {d['estudiante']['nombre']}"),
              p(f"Fecha de entrega: {date.fromisoformat(d['fecha']).strftime('%d/%m/%Y')}"), p("Breve explicación", "h")]
    story.extend(p(t) for t in d["explicacion"])
    if d["estado"] == "entrega_con_observaciones":
        story.append(p("Observación documentada", "h"))
        story.extend(p(t) for t in parrafos_observaciones(d))

    def bloques(posicion):
        for b in d.get("secciones", []):
            if b.get("posicion", "despues_evidencia") != posicion:
                continue
            story.append(p(b["titulo"], "h"))
            story.extend(p(t) for t in b["parrafos"])
            if "tabla" in b:
                t = b["tabla"]
                contenido = [[p(c, "cell_h") for c in t["columnas"]]] + [[p(c, "cell") for c in f] for f in t["filas"]]
                tabla = Table(contenido, colWidths=t["anchos"], repeatRows=1, hAlign="LEFT")
                tabla.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AZUL)),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5FA")]),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]))
                story.extend([tabla, Spacer(1, 8)])

    bloques("antes_evidencia")
    if not d["evidencias"]:
        story.append(p("Evidencia fotográfica", "h"))
        story.append(p("Pendiente de incorporar capturas verificadas de las ejecuciones.", "small"))
    for n, e in enumerate(d["evidencias"], 1):
        w, h = dimensiones(e["_imagen"])
        img = Image(str(e["_imagen"]), width=w, height=h)
        img.hAlign = "CENTER"
        heading = [p("Evidencia fotográfica", "h")] if n == 1 else []
        story.append(KeepTogether(heading + [p(etiqueta_evidencia(e, n), "caption"), img, Spacer(1, 4), p(e["descripcion"], "small")]))
    bloques("despues_evidencia")
    if d["reflexion"]:
        story.append(p("Reflexión", "h"))
        story.extend(p(t) for t in d["reflexion"])
    url = d["repositorio"].get("url", "")
    if mostrar_repositorio(d):
        story.append(p("Enlace del repositorio", "h"))
    if mostrar_repositorio(d) and url:
        story.append(Paragraph(f'<link href="{escape(url)}" color="#{AZUL}">{escape(url)}</link>', styles["body"]))
    elif mostrar_repositorio(d):
        story.append(p("Pendiente de publicar y comprobar el acceso del docente.", "small"))
    if mostrar_repositorio(d) and d["repositorio"].get("acceso_docente_verificado") is True:
        story.append(p("Acceso del docente verificado.", "small"))
    elif url:
        story.append(p("Acceso del docente pendiente de verificación.", "small"))

    def pie(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#ACB7D9"))
        canvas.setLineWidth(0.4)
        canvas.line(52.8, 45, 559.2, 45)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#" + GRIS))
        canvas.drawString(52.8, 32, "CEDIA - Training Management Office | Informe en formato PDF para el AVAC")
        canvas.drawRightString(559.2, 32, str(doc.page))
        canvas.restoreState()

    SimpleDocTemplate(salida, pagesize=(612, 792), leftMargin=52.8, rightMargin=52.8,
                      topMargin=42, bottomMargin=60, title="Informe de Entrega", author=d["estudiante"]["nombre"]).build(story, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()


def docx_bytes(d: dict) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(612), Pt(792)
    sec.left_margin = sec.right_margin = Pt(52.8)
    sec.top_margin, sec.bottom_margin, sec.footer_distance = Pt(42), Pt(60), Pt(25)
    for nombre, tam, color in (("Normal", 10.5, "000000"), ("Title", 20, AZUL), ("Subtitle", 11, GRIS), ("Heading 1", 13, AZUL), ("Caption", 9, GRIS)):
        estilo = doc.styles[nombre]
        # Sustituto métrico libre de Helvetica disponible en el runtime bundled.
        estilo.font.name = "Liberation Sans"
        estilo.font.size = Pt(tam)
        estilo.font.color.rgb = RGBColor.from_string(color)
        estilo.font.bold = nombre in ("Title", "Heading 1")
        estilo.font.italic = nombre == "Caption"
        rpr = estilo.element.get_or_add_rPr()
        fuentes = rpr.find(qn("w:rFonts"))
        for atributo in list(fuentes.attrib):
            if "theme" in atributo.lower():
                del fuentes.attrib[atributo]
        for atributo in ("eastAsia", "cs"):
            fuentes.set(qn(f"w:{atributo}"), "Liberation Sans")
        for tag in ("spacing", "kern", "szCs", "iCs"):
            for nodo in rpr.findall(qn(f"w:{tag}")):
                rpr.remove(nodo)
        ppr = estilo.element.get_or_add_pPr()
        for tag in ("pBdr", "numPr"):
            for nodo in ppr.findall(qn(f"w:{tag}")):
                ppr.remove(nodo)
        estilo.paragraph_format.space_after = Pt(7)
        estilo.paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Caption"].paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.keep_with_next = True
    doc.styles["Subtitle"].paragraph_format.keep_with_next = True
    doc.core_properties.author = d["estudiante"]["nombre"]
    doc.core_properties.title = "Informe de Entrega"
    doc.core_properties.subject = subtitulo(d)
    doc.add_paragraph("Informe de Entrega", "Title")
    sub = doc.add_paragraph(subtitulo(d), "Subtitle")
    borde = OxmlElement("w:pBdr")
    inferior = OxmlElement("w:bottom")
    for k, v in {"val": "single", "sz": "8", "space": "10", "color": AZUL}.items():
        inferior.set(qn(f"w:{k}"), v)
    borde.append(inferior)
    sub._p.get_or_add_pPr().append(borde)
    if d["estado"] in AVISOS_ESTADO:
        doc.add_paragraph(AVISOS_ESTADO[d["estado"]], "Heading 1")
    doc.add_heading("Datos del estudiante", 1)
    for k, v in (("Nombre completo", d["estudiante"]["nombre"]), ("Fecha de entrega", date.fromisoformat(d["fecha"]).strftime("%d/%m/%Y"))):
        p = doc.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        p.add_run(v)
    doc.add_heading("Breve explicación", 1)
    for t in d["explicacion"]:
        doc.add_paragraph(t)
    if d["estado"] == "entrega_con_observaciones":
        doc.add_heading("Observación documentada", 1)
        for t in parrafos_observaciones(d):
            doc.add_paragraph(t)

    def bloques(posicion):
        for b in d.get("secciones", []):
            if b.get("posicion", "despues_evidencia") != posicion:
                continue
            doc.add_heading(b["titulo"], 1)
            for t in b["parrafos"]:
                doc.add_paragraph(t)
            if "tabla" not in b:
                continue
            t = b["tabla"]
            tabla = doc.add_table(rows=1, cols=len(t["columnas"]))
            tabla.autofit = False
            for col, width in zip(tabla.columns, t["anchos"]):
                col.width = Pt(width)
            for i, fila in enumerate([t["columnas"], *t["filas"]]):
                row = tabla.rows[0] if i == 0 else tabla.add_row()
                trpr = row._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if i == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                for j, valor in enumerate(fila):
                    cell = row.cells[j]
                    cell.width = Pt(t["anchos"][j])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = valor
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for side in ("top", "left", "bottom", "right"):
                        border = OxmlElement(f"w:{side}")
                        for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                            border.set(qn(f"w:{k}"), v)
                        borders.append(border)
                    props.append(borders)
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), AZUL if i == 0 else ("F3F5FA" if i % 2 == 0 else "FFFFFF"))
                    props.append(fill)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        mar = OxmlElement(f"w:{side}")
                        mar.set(qn("w:w"), "120")
                        mar.set(qn("w:type"), "dxa")
                        margins.append(mar)
                    props.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(0)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.bold = i == 0
                            run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "000000")
            doc.add_paragraph()

    bloques("antes_evidencia")
    doc.add_heading("Evidencia fotográfica", 1)
    if not d["evidencias"]:
        doc.add_paragraph("Pendiente de incorporar capturas verificadas de las ejecuciones.")
    for n, e in enumerate(d["evidencias"], 1):
        doc.add_paragraph(etiqueta_evidencia(e, n), "Caption")
        w, h = dimensiones(e["_imagen"])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        inline = p.add_run().add_picture(str(e["_imagen"]), width=Pt(w), height=Pt(h))
        inline._inline.docPr.set("descr", e["descripcion"])
        p = doc.add_paragraph(e["descripcion"])
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(GRIS)
    bloques("despues_evidencia")
    if d["reflexion"]:
        doc.add_heading("Reflexión", 1)
        for t in d["reflexion"]:
            doc.add_paragraph(t)
    url = d["repositorio"].get("url", "")
    if mostrar_repositorio(d):
        doc.add_heading("Enlace del repositorio", 1)
    if mostrar_repositorio(d) and url:
        p = doc.add_paragraph()
        enlace = OxmlElement("w:hyperlink")
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        enlace.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = url
        run.append(t)
        enlace.append(run)
        p._p.append(enlace)
    elif mostrar_repositorio(d):
        doc.add_paragraph("Pendiente de publicar y comprobar el acceso del docente.")
    if mostrar_repositorio(d) and d["repositorio"].get("acceso_docente_verificado") is True:
        doc.add_paragraph("Acceso del docente verificado.")
    elif url:
        doc.add_paragraph("Acceso del docente pendiente de verificación.")
    pie = sec.footer.paragraphs[0]
    pie.add_run("CEDIA - Training Management Office | Informe en formato PDF para el AVAC   ")
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    pie._p.append(campo)
    for r in pie.runs:
        r.font.name = "Liberation Sans"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(GRIS)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def generar(d: dict, salida: Path) -> dict:
    """Valida/renderiza en memoria antes de reservar una carpeta de salida nueva."""
    if salida.exists():
        raise ValueError("La carpeta de salida ya existe; usar una ruta nueva para preservar versiones")
    pdf, docx = pdf_bytes(d), docx_bytes(d)
    for e in d["evidencias"]:
        if sha256(e["_imagen"]) != e["_sha256_imagen"] or sha256(e["_fuente"]) != e["_sha256_fuente"]:
            raise ValueError(f"La evidencia {e['id']} cambió durante la generación; salida preservada sin crear")
    evidencia = [{"id": e["id"], "tipo": e["tipo"], "imagen": e["_imagen"].name,
                  "fuente": e["_fuente"].name, "sha256_imagen": e["_sha256_imagen"],
                  "sha256_fuente": e["_sha256_fuente"], "verificada": e.get("verificada", False)} for e in d["evidencias"]]
    nombre = f"Informe-Practica-{d['practica']['numero']:02d}"
    archivos = {nombre + ".pdf": pdf, nombre + ".docx": docx}
    manifiesto = {"version": 1, "estado": d["estado"], "practica": d["practica"]["numero"],
                  "generado": datetime.now(ZoneInfo("America/Guayaquil")).isoformat(),
                  "entrada_sha256": sha256(d["_entrada"]),
                  "artefactos": {n: {"bytes": len(b), "sha256": hashlib.sha256(b).hexdigest()} for n, b in archivos.items()},
                  "evidencias": evidencia,
                  "limite_verificacion": "Los hashes verifican integridad; la procedencia y las afirmaciones deben comprobarse contra la ejecución original."}
    if d["estado"] == "entrega_con_observaciones":
        manifiesto.update(cumplimiento_total=False, requisitos_pendientes=["manejo_429"],
                          observaciones=d["observaciones"])
    plantilla = Path(__file__).resolve().parent.parent / "materiales" / "Plantilla-Informe.pdf"
    if plantilla.is_file():
        manifiesto["plantilla_sha256"] = sha256(plantilla)
    salida.mkdir(parents=True, exist_ok=False)
    for nombre_archivo, contenido in archivos.items():
        with (salida / nombre_archivo).open("xb") as f:
            f.write(contenido)
    with (salida / "manifiesto.json").open("x", encoding="utf-8") as f:
        json.dump(manifiesto, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return manifiesto


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("datos", type=Path)
    parser.add_argument("--salida", type=Path, help="Carpeta nueva obligatoria para generar")
    parser.add_argument("--validar", action="store_true", help="Validar sin escribir archivos")
    args = parser.parse_args()
    try:
        datos = cargar(args.datos.resolve())
        if args.validar:
            resultado = {"ok": True, "estado": datos["estado"], "evidencias": len(datos["evidencias"])}
            if datos["estado"] == "entrega_con_observaciones":
                resultado.update(cumplimiento_total=False, requisitos_pendientes=["manejo_429"])
            print(json.dumps(resultado))
        elif args.salida:
            print(json.dumps(generar(datos, args.salida.resolve()), ensure_ascii=False, indent=2))
        else:
            parser.error("se necesita --salida o --validar")
    except (ValueError, OSError, KeyError, TypeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
